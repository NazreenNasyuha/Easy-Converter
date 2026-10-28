"""Unit and integration test suite for Easy Converter across all 9 file categories and daily tools."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from PySide6.QtCore import QPointF, Qt
from PySide6.QtGui import QGuiApplication
from PySide6.QtWidgets import QApplication

from core.action_registry import ModifierMode, format_center_badge, get_actions_for_files
from core.converters import execute_action
from core.radial_overlay import RadialOverlay
from core.ui.image_editor_dialog import ImageEditorDialog
from core.ui.radial_math import cartesian_to_polar_cw, hit_test_wedge


class EasyConverterTestSuite(unittest.TestCase):
    """Comprehensive verification suite for geometry, registry, and all 9 conversion domains."""

    @classmethod
    def setUpClass(cls) -> None:
        QGuiApplication.setHighDpiScaleFactorRoundingPolicy(
            Qt.HighDpiScaleFactorRoundingPolicy.PassThrough
        )
        cls.app = QApplication.instance() or QApplication(sys.argv)

    def test_polar_math_and_wedge_hit_testing(self) -> None:
        r, deg = cartesian_to_polar_cw(190.0, 90.0, 190.0, 190.0)
        self.assertAlmostEqual(r, 100.0)
        self.assertAlmostEqual(deg, 0.0)

        _, deg_right = cartesian_to_polar_cw(290.0, 190.0, 190.0, 190.0)
        self.assertAlmostEqual(deg_right, 90.0)

        center_hit, idx = hit_test_wedge(12.0, 0.0, 5, 48.0, 170.0)
        self.assertTrue(center_hit)
        self.assertEqual(idx, -1)

        center_hit, idx = hit_test_wedge(100.0, 10.0, 5, 48.0, 170.0)
        self.assertFalse(center_hit)
        self.assertEqual(idx, 0)

    def test_all_9_categories_have_format_and_tools_slices(self) -> None:
        sample_files_by_category = {
            "Images": "photo.png",
            "Video": "clip.mp4",
            "Audio": "song.mp3",
            "PDF": "report.pdf",
            "Document": "notes.docx",
            "Slides": "deck.pptx",
            "Spreadsheet": "table.xlsx",
            "Archives": "bundle.zip",
            "Drawing": "blueprint.dxf",
        }
        for category, filename in sample_files_by_category.items():
            ext = Path(filename).suffix.lower()
            shift_slices = get_actions_for_files([filename], ModifierMode.SHIFT)
            alt_slices = get_actions_for_files([filename], ModifierMode.ALT)
            ctrl_slices = get_actions_for_files([filename], ModifierMode.CTRL)

            self.assertTrue(
                3 <= len(shift_slices) <= 8,
                f"{category} Shift slices out of range: {len(shift_slices)}",
            )
            self.assertTrue(
                6 <= len(alt_slices) <= 8,
                f"{category} Alt slices should have 6-8 daily tools, got {len(alt_slices)}",
            )
            self.assertEqual(len(ctrl_slices), 5)
            self.assertTrue(
                all(s.target_ext != ext for s in shift_slices),
                f"{category} Shift mode failed to filter source extension {ext}",
            )

        self.assertEqual(format_center_badge(["photo.heic"]), ".HEIC")
        self.assertEqual(format_center_badge(["a.png", "b.png", "c.png"]), "3 Files")

    def test_images_remove_bg_crop_upscale_studio_and_pdf_tools(self) -> None:
        with tempfile.TemporaryDirectory(prefix="easy_converter_img_") as tmp:
            tmp_dir = Path(tmp)
            sample_png = tmp_dir / "sample.png"

            # Create image with uniform border and central subject
            img = Image.new("RGBA", (240, 180), (245, 245, 245, 255))
            draw = ImageDraw.Draw(img)
            draw.ellipse((60, 40, 180, 140), fill=(255, 140, 66, 255))
            img.save(sample_png, "PNG")

            for action_id, label, expected_ext in [
                ("img_to_webp", "WEBP", ".webp"),
                ("img_to_jpg", "JPG", ".jpg"),
                ("img_to_ico", "ICO", ".ico"),
                ("img_to_pdf", "PDF", ".pdf"),
                ("img_remove_bg", "Remove BG", ".png"),
                ("img_crop_edit", "Crop & Edit", ".png"),
                ("img_crop_square", "Square 1:1", ".png"),
                ("img_upscale_2x", "Upscale 2x", ".png"),
                ("img_auto_enhance", "Auto Enhance", ".png"),
            ]:
                res = execute_action(action_id, label, [str(sample_png)])
                self.assertTrue(res.success, f"{action_id} failed: {res.error_message}")
                out_file = Path(res.output_paths[0])
                self.assertTrue(out_file.exists() and out_file.stat().st_size > 0)
                self.assertEqual(out_file.suffix.lower(), expected_ext)

            # Verify ImageEditorDialog initializes, applies crop/BG removal, and saves
            studio = ImageEditorDialog(sample_png)
            studio._on_aspect_changed("1:1 Square")
            studio._on_toggle_remove_bg()
            studio._on_rotate()
            saved_paths: list[str] = []
            studio.image_saved.connect(saved_paths.append)
            studio._on_save()
            self.assertEqual(len(saved_paths), 1)
            self.assertTrue(Path(saved_paths[0]).exists())

            # PDF Daily Tools (autocrop, rotate 90, page numbers, grayscale, strip meta)
            sample_pdf = tmp_dir / "sample.pdf"
            for pdf_act in (
                "pdf_to_png",
                "pdf_to_docx",
                "pdf_to_svg",
                "pdf_autocrop",
                "pdf_rotate_90",
                "pdf_add_page_numbers",
                "pdf_grayscale",
                "pdf_strip_meta",
            ):
                res_pdf = execute_action(pdf_act, pdf_act, [str(sample_pdf)])
                self.assertTrue(res_pdf.success, f"{pdf_act} failed: {res_pdf.error_message}")
                self.assertTrue(Path(res_pdf.output_paths[0]).exists())

    def test_document_slides_spreadsheet_archive_and_drawing_daily_tools(self) -> None:
        with tempfile.TemporaryDirectory(prefix="easy_converter_domains_") as tmp:
            tmp_dir = Path(tmp)

            # 1. Document Tools
            sample_md = tmp_dir / "readme.md"
            sample_md.write_text(
                "# Project Title\n\nContact admin@example.com or visit https://example.com.\n"
                "Call 555-123-4567.\n\n## Section\nMore text.",
                encoding="utf-8",
            )
            for act in (
                "doc_to_docx",
                "doc_to_pdf",
                "doc_to_html",
                "doc_clean_text",
                "doc_extract_links",
                "doc_extract_outline",
                "doc_redact_pii",
                "doc_word_count",
            ):
                res = execute_action(act, act, [str(sample_md)])
                self.assertTrue(res.success, f"{act} failed: {res.error_message}")

            # 2. Slides Tools
            from pptx import Presentation

            sample_pptx = tmp_dir / "pitch.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            if slide.shapes.title:
                slide.shapes.title.text = "Quarterly Review https://example.com"
            prs.save(str(sample_pptx))

            for act in (
                "slides_to_pdf",
                "slides_to_png",
                "slides_to_md",
                "slides_handout_pdf",
                "slides_export_notes",
                "slides_strip_notes",
                "slides_extract_links",
                "slides_stats",
            ):
                res = execute_action(act, act, [str(sample_pptx)])
                self.assertTrue(res.success, f"{act} failed: {res.error_message}")

            # 3. Spreadsheet Tools
            sample_csv = tmp_dir / "metrics.csv"
            sample_csv.write_text("name,score\nCharlie,82\nAlice,95\nBob,88\nAlice,95\n", encoding="utf-8")
            for act in (
                "sheet_to_xlsx",
                "sheet_to_json",
                "sheet_to_md",
                "sheet_to_pdf",
                "sheet_deduplicate",
                "sheet_transpose",
                "sheet_sort_asc",
                "sheet_autofit_xlsx",
                "sheet_column_stats",
                "sheet_stats",
            ):
                res = execute_action(act, act, [str(sample_csv)])
                self.assertTrue(res.success, f"{act} failed: {res.error_message}")

            # 4. Archive Tools
            res_zip = execute_action("pack_zip", "Pack .ZIP", [str(sample_md), str(sample_csv)])
            self.assertTrue(res_zip.success)
            created_zip = Path(res_zip.output_paths[0])
            for act in (
                "arch_to_targz",
                "arch_list_contents",
                "arch_verify",
                "arch_clean_junk",
                "arch_flatten_extract",
                "arch_extract_folder",
            ):
                res = execute_action(act, act, [str(created_zip)])
                self.assertTrue(res.success, f"{act} failed: {res.error_message}")

            # 5. Drawing Tools
            sample_svg = tmp_dir / "diagram.svg"
            sample_svg.write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="200">'
                '<line x1="10" y1="10" x2="190" y2="190" stroke="black"/>'
                '<circle cx="100" cy="100" r="40" stroke="red" fill="none"/></svg>',
                encoding="utf-8",
            )
            res_dxf = execute_action("draw_to_dxf", "DXF", [str(sample_svg)])
            self.assertTrue(res_dxf.success, res_dxf.error_message)
            created_dxf = Path(res_dxf.output_paths[0])

            for act in (
                "draw_to_png",
                "draw_transparent_png",
                "draw_blueprint_pdf",
                "draw_extract_layers",
                "draw_to_pdf",
                "draw_to_svg",
                "draw_inspect_info",
            ):
                res = execute_action(act, act, [str(created_dxf)])
                self.assertTrue(res.success, f"{act} on DXF failed: {res.error_message}")

    def test_radial_overlay_widget_lifecycle(self) -> None:
        overlay = RadialOverlay()
        overlay.open_demo_mode()
        overlay._update_hit_test(QPointF(190.0, 190.0))
        self.assertTrue(overlay._center_hovered)
        self.assertEqual(overlay._hovered_index, -1)

        overlay._update_hit_test(QPointF(190.0, 90.0))
        self.assertFalse(overlay._center_hovered)
        self.assertEqual(overlay._hovered_index, 0)

        overlay.set_modifier_mode(ModifierMode.ALT)
        self.assertEqual(overlay._mode, ModifierMode.ALT)

        overlay.close_overlay()
        self.assertFalse(overlay.isVisible())


def run_cli_self_test() -> int:
    """Run the test suite and print a clean summary."""
    print("=== Easy Converter Automated Test Suite (9 Categories + Daily Tools) ===")
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(EasyConverterTestSuite)
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    if result.wasSuccessful():
        print("=== ALL SELF-TESTS PASSED ===")
        return 0
    return 1
