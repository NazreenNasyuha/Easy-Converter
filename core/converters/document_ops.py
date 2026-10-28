"""Document conversion and daily text/privacy/outline tools (.docx, .txt, .md, .html, .rtf, .epub)."""

from __future__ import annotations

import html
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from core.converters.path_utils import format_saved_summary, get_safe_output_path
from core.models import ConversionResult

try:
    import pymupdf as fitz
except ImportError:
    import fitz


DOCUMENT_ACTION_IDS: frozenset[str] = frozenset({
    # Document Formats (Shift Mode)
    "doc_to_pdf",
    "doc_to_docx",
    "doc_to_txt",
    "doc_to_md",
    "doc_to_html",
    # Document Daily Tools (Alt Mode)
    "doc_clean_text",
    "doc_extract_images",
    "doc_extract_links",
    "doc_extract_outline",
    "doc_redact_pii",
    "doc_word_count",
    "doc_to_pdf_bw",
})


class _HTMLTextExtractor(HTMLParser):
    """Lightweight HTML-to-text extractor preserving block breaks."""

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n")
        if tag == "li":
            self.parts.append("• ")

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    def get_text(self) -> str:
        raw = "".join(self.parts)
        return re.sub(r"\n{3,}", "\n\n", raw).strip()


def extract_document_paragraphs(src: Path) -> list[str]:
    """Extract ordered paragraphs of text from a document file."""
    ext = src.suffix.lower()

    if ext == ".docx":
        from docx import Document

        doc = Document(str(src))
        paras: list[str] = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            style_name = (p.style.name if p.style else "").lower()
            if "heading 1" in style_name:
                paras.append(f"# {t}")
            elif "heading 2" in style_name:
                paras.append(f"## {t}")
            elif "heading 3" in style_name:
                paras.append(f"### {t}")
            else:
                paras.append(t)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    paras.append(row_text)
        return paras

    if ext == ".epub":
        paras = []
        with fitz.open(str(src)) as book:
            for page in book:
                text = page.get_text("text").strip()
                if text:
                    paras.extend(p.strip() for p in text.split("\n\n") if p.strip())
        return paras

    raw_text = src.read_text(encoding="utf-8", errors="replace")

    if ext in (".html", ".htm"):
        parser = _HTMLTextExtractor()
        parser.feed(raw_text)
        text = parser.get_text()
        return [p.strip() for p in text.split("\n") if p.strip()]

    if ext == ".rtf":
        stripped = re.sub(r"\\par[d]?\b", "\n", raw_text)
        stripped = re.sub(r"\\[a-zA-Z]+-?\d*\s?", "", stripped)
        stripped = stripped.replace("{", "").replace("}", "")
        return [line.strip() for line in stripped.splitlines() if line.strip()]

    return [line.rstrip() for line in raw_text.splitlines()]


def render_text_paragraphs_to_pdf(
    paragraphs: list[str],
    dst: Path,
    title: str = "",
    grayscale: bool = False,
) -> Path:
    """Render paragraphs cleanly onto paginated A4 PDF pages using PyMuPDF."""
    doc = fitz.open()
    page_w, page_h = 595.0, 842.0  # A4 points
    margin = 50.0
    usable_w = page_w - 2 * margin
    y = margin

    page = doc.new_page(width=page_w, height=page_h)
    title_color = (0.15, 0.15, 0.15) if grayscale else (0.12, 0.16, 0.23)
    body_color = (0.1, 0.1, 0.1)

    if title:
        page.insert_text(
            fitz.Point(margin, y + 14),
            title,
            fontsize=16,
            fontname="helv",
            color=title_color,
        )
        y += 32.0

    for para in paragraphs:
        text = para.strip()
        if not text:
            y += 10.0
            continue

        is_heading = text.startswith("#")
        clean_line = text.lstrip("#").strip() if is_heading else text
        font_size = 13.0 if is_heading else 10.5
        line_height = font_size * 1.45

        chars_per_line = max(25, int(usable_w / (font_size * 0.52)))
        words = clean_line.split()
        lines: list[str] = []
        current = ""
        for w in words:
            cand = f"{current} {w}".strip() if current else w
            if len(cand) <= chars_per_line:
                current = cand
            else:
                if current:
                    lines.append(current)
                current = w
        if current:
            lines.append(current)

        block_h = max(1, len(lines)) * line_height + 6.0
        if y + block_h > page_h - margin:
            page = doc.new_page(width=page_w, height=page_h)
            y = margin

        for ln in lines:
            page.insert_text(
                fitz.Point(margin, y + font_size),
                ln,
                fontsize=font_size,
                fontname="helv",
                color=title_color if is_heading else body_color,
            )
            y += line_height
        y += 4.0

    doc.save(str(dst), garbage=4, deflate=True)
    doc.close()
    return dst


def run_document_action(
    action_id: str,
    action_label: str,
    src: Path,
) -> ConversionResult:
    """Execute a document format conversion or daily text/privacy tool operation."""
    paragraphs = extract_document_paragraphs(src)
    full_text = "\n\n".join(p for p in paragraphs if p.strip())

    if action_id == "doc_word_count":
        words = full_text.split()
        chars = len(full_text)
        read_min = max(1, round(len(words) / 200))
        summary = f"{len(words)} Words | {chars} Chars (~{read_min} min read)"
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(src)],
            clipboard_text=f"{src.name}\n{summary}",
            summary_text=summary,
        )

    if action_id in ("doc_to_pdf", "doc_to_pdf_bw"):
        is_bw = action_id == "doc_to_pdf_bw"
        dst = get_safe_output_path(src, new_ext=".pdf", suffix="_bw" if is_bw else "")
        render_text_paragraphs_to_pdf(paragraphs, dst, title=src.stem, grayscale=is_bw)
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "doc_to_docx":
        from docx import Document

        dst = get_safe_output_path(src, new_ext=".docx")
        doc = Document()
        doc.add_heading(src.stem, level=1)
        for p in paragraphs:
            stripped = p.strip()
            if not stripped:
                continue
            if stripped.startswith("# "):
                doc.add_heading(stripped[2:].strip(), level=1)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:].strip(), level=2)
            elif stripped.startswith("### "):
                doc.add_heading(stripped[4:].strip(), level=3)
            else:
                doc.add_paragraph(stripped)
        doc.save(str(dst))
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "doc_to_txt":
        dst = get_safe_output_path(src, new_ext=".txt")
        dst.write_text(full_text, encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "doc_to_md":
        dst = get_safe_output_path(src, new_ext=".md")
        md_lines = [f"# {src.stem}", ""]
        for p in paragraphs:
            if p.strip():
                md_lines.append(p.strip())
                md_lines.append("")
        dst.write_text("\n".join(md_lines).strip() + "\n", encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "doc_to_html":
        dst = get_safe_output_path(src, new_ext=".html")
        body_tags: list[str] = [f"<h1>{html.escape(src.stem)}</h1>"]
        for p in paragraphs:
            s = p.strip()
            if not s:
                continue
            if s.startswith("# "):
                body_tags.append(f"<h2>{html.escape(s[2:].strip())}</h2>")
            elif s.startswith("## "):
                body_tags.append(f"<h3>{html.escape(s[3:].strip())}</h3>")
            else:
                body_tags.append(f"<p>{html.escape(s)}</p>")
        html_doc = (
            "<!DOCTYPE html>\n<html lang='en'>\n<head>\n<meta charset='utf-8'>\n"
            f"<title>{html.escape(src.stem)}</title>\n"
            "<style>body{font-family:system-ui,sans-serif;max-width:780px;margin:2rem auto;"
            "line-height:1.6;color:#1e293b;padding:0 1rem;}</style>\n"
            "</head>\n<body>\n" + "\n".join(body_tags) + "\n</body>\n</html>\n"
        )
        dst.write_text(html_doc, encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    # --- Alt Mode: Daily Document Tools ---
    if action_id == "doc_clean_text":
        out_ext = src.suffix if src.suffix.lower() in (".txt", ".md") else ".txt"
        dst = get_safe_output_path(src, new_ext=out_ext, suffix="_clean")
        cleaned = [re.sub(r"[ \t]+", " ", p).strip() for p in paragraphs if p.strip()]
        dst.write_text("\n\n".join(cleaned) + "\n", encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "doc_extract_links":
        urls = sorted(set(re.findall(r"https?://[^\s<>\"')]+", full_text)))
        emails = sorted(set(re.findall(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", full_text)))
        dst = get_safe_output_path(src, new_ext=".txt", suffix="_links")
        lines = [
            f"=== Extracted Links & Emails: {src.name} ===",
            f"URLs ({len(urls)}):",
            *(urls if urls else ["  (None)"]),
            "",
            f"Emails ({len(emails)}):",
            *(emails if emails else ["  (None)"]),
        ]
        report = "\n".join(lines) + "\n"
        dst.write_text(report, encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            clipboard_text="\n".join([*urls, *emails]) if (urls or emails) else report,
            summary_text=f"Found {len(urls)} URLs & {len(emails)} emails → {dst.name}",
        )

    if action_id == "doc_extract_outline":
        dst = get_safe_output_path(src, new_ext=".md", suffix="_outline")
        headings: list[str] = [f"# Outline: {src.stem}", ""]
        for p in paragraphs:
            s = p.strip()
            if s.startswith("#") or (len(s) < 85 and s.endswith(":")) or s.isupper():
                headings.append(s if s.startswith("#") else f"- **{s}**")
        if len(headings) <= 2:
            # Summarize first sentence of each paragraph
            for p in paragraphs[:25]:
                s = p.strip()
                if s:
                    first_sentence = re.split(r"(?<=[.!?])\s+", s)[0]
                    headings.append(f"- {first_sentence[:110]}")
        dst.write_text("\n".join(headings) + "\n", encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "doc_redact_pii":
        out_ext = src.suffix if src.suffix.lower() in (".txt", ".md") else ".txt"
        dst = get_safe_output_path(src, new_ext=out_ext, suffix="_redacted")
        redacted = full_text
        # Mask emails
        redacted = re.sub(
            r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+",
            "[REDACTED_EMAIL]",
            redacted,
        )
        # Mask phone numbers
        redacted = re.sub(
            r"\b(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]\d{3}[-.\s]\d{4}\b",
            "[REDACTED_PHONE]",
            redacted,
        )
        # Mask SSN / 16-digit card numbers
        redacted = re.sub(r"\b\d{3}-\d{2}-\d{4}\b", "[REDACTED_ID]", redacted)
        redacted = re.sub(r"\b(?:\d{4}[-\s]?){3}\d{4}\b", "[REDACTED_CARD]", redacted)
        dst.write_text(redacted + "\n", encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "doc_extract_images":
        media_dir = src.parent / f"{src.stem}_media"
        extracted: list[Path] = []
        if zipfile.is_zipfile(src):
            with zipfile.ZipFile(src, "r") as zf:
                for member in zf.namelist():
                    if member.startswith("word/media/") or member.lower().endswith(
                        (".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg")
                    ):
                        media_dir.mkdir(parents=True, exist_ok=True)
                        out_file = get_safe_output_path(media_dir / Path(member).name)
                        out_file.write_bytes(zf.read(member))
                        extracted.append(out_file)
        if not extracted:
            return ConversionResult(
                success=True,
                action_id=action_id,
                action_label=action_label,
                output_paths=[str(src)],
                summary_text="No embedded images found in document",
            )
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(p) for p in extracted],
            summary_text=f"Extracted {len(extracted)} image{'s' if len(extracted) > 1 else ''}",
        )

    raise ValueError(f"Unknown document action_id: {action_id}")
