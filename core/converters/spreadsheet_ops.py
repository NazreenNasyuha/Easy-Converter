"""Spreadsheet conversion and daily tabular data tools (.xlsx, .csv, .tsv, .json)."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from core.converters.document_ops import render_text_paragraphs_to_pdf
from core.converters.path_utils import format_saved_summary, get_safe_output_path
from core.models import ConversionResult

SPREADSHEET_ACTION_IDS: frozenset[str] = frozenset({
    # Spreadsheet Formats (Shift Mode)
    "sheet_to_xlsx",
    "sheet_to_csv",
    "sheet_to_tsv",
    "sheet_to_json",
    "sheet_to_pdf",
    "sheet_to_md",
    # Spreadsheet Daily Tools (Alt Mode)
    "sheet_deduplicate",
    "sheet_trim_clean",
    "sheet_transpose",
    "sheet_sort_asc",
    "sheet_autofit_xlsx",
    "sheet_split_sheets",
    "sheet_column_stats",
    "sheet_stats",
})


def load_tabular_sheets(src: Path) -> dict[str, list[list[str]]]:
    """Load one or more sheets of rows (`list[list[str]]`) from `.xlsx`, `.csv`, `.tsv`, or `.json`."""
    ext = src.suffix.lower()

    if ext in (".xlsx", ".xls"):
        from openpyxl import load_workbook

        wb = load_workbook(filename=str(src), data_only=True)
        sheets: dict[str, list[list[str]]] = {}
        for name in wb.sheetnames:
            ws = wb[name]
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                str_row = ["" if v is None else str(v) for v in row]
                if any(cell.strip() for cell in str_row):
                    rows.append(str_row)
            sheets[name] = rows
        wb.close()
        return sheets or {"Sheet1": []}

    raw_text = src.read_text(encoding="utf-8-sig", errors="replace")

    if ext == ".json":
        data = json.loads(raw_text)
        if isinstance(data, list) and data:
            if isinstance(data[0], dict):
                headers = list(data[0].keys())
                rows = [headers]
                for item in data:
                    if isinstance(item, dict):
                        rows.append(["" if item.get(k) is None else str(item.get(k)) for k in headers])
                return {src.stem: rows}
            if isinstance(data[0], list):
                return {src.stem: [[str(c) for c in r] for r in data]}
        return {src.stem: [[str(data)]]}

    delimiter = "\t" if ext == ".tsv" else ","
    reader = csv.reader(raw_text.splitlines(), delimiter=delimiter)
    rows = [list(r) for r in reader if any(c.strip() for c in r)]
    return {src.stem: rows}


def _write_delimited(rows: list[list[str]], dst: Path, delimiter: str) -> Path:
    with dst.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter=delimiter)
        writer.writerows(rows)
    return dst


def _write_xlsx(
    sheets: dict[str, list[list[str]]],
    dst: Path,
    autofit: bool = False,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    default_ws = wb.active
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, (sheet_name, rows) in enumerate(sheets.items()):
        safe_title = (sheet_name or f"Sheet{idx + 1}")[:31]
        ws = default_ws if idx == 0 and default_ws is not None else wb.create_sheet(title=safe_title)
        ws.title = safe_title
        for r_idx, row in enumerate(rows, start=1):
            ws.append(row)
            if r_idx == 1:
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF" if autofit else "000000")
                    if autofit:
                        cell.fill = header_fill
            elif autofit and r_idx % 2 == 0:
                for cell in ws[r_idx]:
                    cell.fill = zebra_fill

        if autofit:
            ws.freeze_panes = "A2"
            for col in ws.columns:
                max_len = max((len(str(c.value or "")) for c in col), default=8)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 10), 48)

    wb.save(str(dst))
    wb.close()
    return dst


def run_spreadsheet_action(
    action_id: str,
    action_label: str,
    src: Path,
) -> ConversionResult:
    """Execute a spreadsheet format conversion or daily tabular tool operation."""
    sheets = load_tabular_sheets(src)
    first_sheet_name = next(iter(sheets))
    rows = sheets[first_sheet_name]

    if action_id == "sheet_stats":
        row_count = len(rows)
        col_count = max((len(r) for r in rows), default=0)
        sheet_count = len(sheets)
        summary = f"{row_count} Rows × {col_count} Cols ({sheet_count} Sheet{'s' if sheet_count > 1 else ''})"
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(src)],
            clipboard_text=f"{src.name}: {summary}",
            summary_text=summary,
        )

    if action_id == "sheet_to_xlsx":
        dst = get_safe_output_path(src, new_ext=".xlsx")
        _write_xlsx(sheets, dst, autofit=False)
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_autofit_xlsx":
        dst = get_safe_output_path(src, new_ext=".xlsx", suffix="_styled")
        _write_xlsx(sheets, dst, autofit=True)
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_to_csv":
        dst = get_safe_output_path(src, new_ext=".csv")
        _write_delimited(rows, dst, ",")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_to_tsv":
        dst = get_safe_output_path(src, new_ext=".tsv")
        _write_delimited(rows, dst, "\t")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_to_json":
        dst = get_safe_output_path(src, new_ext=".json")
        if len(rows) >= 2:
            headers = [h.strip() or f"col_{i + 1}" for i, h in enumerate(rows[0])]
            records = [
                {headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers))}
                for r in rows[1:]
            ]
        else:
            records = rows
        dst.write_text(json.dumps(records, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_to_md":
        dst = get_safe_output_path(src, new_ext=".md")
        if not rows:
            dst.write_text(f"# {src.stem}\n\n*(Empty table)*\n", encoding="utf-8")
        else:
            col_n = max(len(r) for r in rows)
            padded = [r + [""] * (col_n - len(r)) for r in rows]
            header_line = "| " + " | ".join(c.replace("|", "\\|") for c in padded[0]) + " |"
            sep_line = "| " + " | ".join("---" for _ in range(col_n)) + " |"
            body_lines = [
                "| " + " | ".join(c.replace("|", "\\|") for c in r) + " |"
                for r in padded[1:]
            ]
            md_text = "\n".join([f"# {src.stem}", "", header_line, sep_line, *body_lines]) + "\n"
            dst.write_text(md_text, encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_to_pdf":
        dst = get_safe_output_path(src, new_ext=".pdf")
        formatted_rows = ["  |  ".join(r) for r in rows]
        render_text_paragraphs_to_pdf(formatted_rows, dst, title=f"{src.stem} ({first_sheet_name})")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_deduplicate":
        seen: set[tuple[str, ...]] = set()
        unique_rows: list[list[str]] = []
        for r in rows:
            key = tuple(c.strip() for c in r)
            if key not in seen:
                seen.add(key)
                unique_rows.append(r)
        out_ext = ".xlsx" if src.suffix.lower() in (".xlsx", ".xls") else ".csv"
        dst = get_safe_output_path(src, new_ext=out_ext, suffix="_dedup")
        if out_ext == ".xlsx":
            _write_xlsx({first_sheet_name: unique_rows}, dst, autofit=True)
        else:
            _write_delimited(unique_rows, dst, ",")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_trim_clean":
        cleaned_rows = [
            [c.strip() for c in r]
            for r in rows
            if any(c.strip() for c in r)
        ]
        out_ext = ".xlsx" if src.suffix.lower() in (".xlsx", ".xls") else ".csv"
        dst = get_safe_output_path(src, new_ext=out_ext, suffix="_clean")
        if out_ext == ".xlsx":
            _write_xlsx({first_sheet_name: cleaned_rows}, dst, autofit=True)
        else:
            _write_delimited(cleaned_rows, dst, ",")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_transpose":
        if not rows:
            transposed: list[list[str]] = []
        else:
            col_n = max(len(r) for r in rows)
            padded = [r + [""] * (col_n - len(r)) for r in rows]
            transposed = [[padded[r_i][c_i] for r_i in range(len(padded))] for c_i in range(col_n)]
        out_ext = ".xlsx" if src.suffix.lower() in (".xlsx", ".xls") else ".csv"
        dst = get_safe_output_path(src, new_ext=out_ext, suffix="_transposed")
        if out_ext == ".xlsx":
            _write_xlsx({first_sheet_name: transposed}, dst, autofit=True)
        else:
            _write_delimited(transposed, dst, ",")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_sort_asc":
        if len(rows) > 2:
            header = rows[0]
            body = sorted(rows[1:], key=lambda r: (r[0].strip().lower() if r else ""))
            sorted_rows = [header, *body]
        else:
            sorted_rows = rows
        out_ext = ".xlsx" if src.suffix.lower() in (".xlsx", ".xls") else ".csv"
        dst = get_safe_output_path(src, new_ext=out_ext, suffix="_sorted")
        if out_ext == ".xlsx":
            _write_xlsx({first_sheet_name: sorted_rows}, dst, autofit=True)
        else:
            _write_delimited(sorted_rows, dst, ",")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_column_stats":
        dst = get_safe_output_path(src, new_ext=".md", suffix="_stats")
        report_lines = [f"# Column Summary Statistics: {src.name}", ""]
        if len(rows) >= 2:
            headers = rows[0]
            report_lines.append("| Column | Count | Numeric Sum | Mean | Min | Max |")
            report_lines.append("| --- | ---: | ---: | ---: | ---: | ---: |")
            for col_i, h_name in enumerate(headers):
                vals = [r[col_i].strip() for r in rows[1:] if col_i < len(r) and r[col_i].strip()]
                nums: list[float] = []
                for v in vals:
                    try:
                        nums.append(float(v.replace(",", "")))
                    except ValueError:
                        pass
                if nums:
                    s_val = sum(nums)
                    m_val = s_val / len(nums)
                    report_lines.append(
                        f"| {h_name or f'Col {col_i+1}'} | {len(vals)} | {s_val:,.2f} | "
                        f"{m_val:,.2f} | {min(nums):,.2f} | {max(nums):,.2f} |"
                    )
                else:
                    report_lines.append(
                        f"| {h_name or f'Col {col_i+1}'} | {len(vals)} | — | — | — | — |"
                    )
        else:
            report_lines.append("*(Not enough rows for column statistics)*")
        dst.write_text("\n".join(report_lines) + "\n", encoding="utf-8")
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(dst)],
            summary_text=format_saved_summary([dst]),
        )

    if action_id == "sheet_split_sheets":
        created: list[Path] = []
        for sheet_name, s_rows in sheets.items():
            safe_sheet = "".join(ch if ch.isalnum() or ch in ("_", "-") else "_" for ch in sheet_name)
            dst = get_safe_output_path(src, new_ext=".csv", suffix=f"_{safe_sheet}")
            _write_delimited(s_rows, dst, ",")
            created.append(dst)
        return ConversionResult(
            success=True,
            action_id=action_id,
            action_label=action_label,
            output_paths=[str(p) for p in created],
            summary_text=format_saved_summary(created),
        )

    raise ValueError(f"Unknown spreadsheet action_id: {action_id}")
